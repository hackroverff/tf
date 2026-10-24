"""One-workbook Excel export for data, analytics, assumptions and approvals."""
from __future__ import annotations

from io import BytesIO
import re

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def excel_export_bytes(tables: dict[str, pd.DataFrame]) -> bytes:
    """Write all tables to a formatted in-memory workbook (no external service)."""
    buffer = BytesIO()
    used_names: set[str] = set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for raw_name, frame in tables.items():
            if frame is None:
                continue
            base = re.sub(r"[^A-Za-z0-9_ -]", "_", str(raw_name))[:31] or "Sheet"
            name = base
            suffix = 2
            while name in used_names:
                name = f"{base[:28]}_{suffix}"
                suffix += 1
            used_names.add(name)
            export = frame.copy()
            # Excel cannot store timezone-aware datetimes; synthetic data is naive,
            # but this makes approval timestamps safe as well.
            for column in export.columns:
                if pd.api.types.is_datetime64tz_dtype(export[column].dtype):
                    export[column] = export[column].dt.tz_localize(None)
            export.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            if len(export.columns) > 0:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17324D")
            for i, column in enumerate(export.columns[:30], start=1):
                sample = export[column].head(150).astype(str)
                width = min(42, max(10, len(str(column)) + 2, int(sample.str.len().quantile(0.9)) + 2 if not sample.empty else 10))
                sheet.column_dimensions[get_column_letter(i)].width = width
    return buffer.getvalue()
