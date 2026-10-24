from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.export import excel_export_bytes
from src.figures import depot_cover_heatmap, ramp_lead_figure


def test_key_new_figures_export_real_pngs() -> None:
    ramp = pd.DataFrame({
        "initial_dp_qc_multiplier": [1.0, 2.5],
        "first_amber": [pd.Timestamp("2026-03-22"), pd.Timestamp("2026-03-18")],
        "first_red": [pd.NaT, pd.Timestamp("2026-03-19")],
        "amber_reaction_lead_days": [9.0, 13.0],
        "red_confirmation_lead_days": [float("nan"), 12.0],
    })
    cover = pd.DataFrame({
        "flavor_id": ["DP", "DP", "MPB", "MPB"],
        "depot_id": ["D01", "D02", "D01", "D02"],
        "cover_days_current_demand": [3.0, 45.0, 21.0, float("inf")],
    })
    for figure in [ramp_lead_figure(ramp), depot_cover_heatmap(cover)]:
        png = figure.to_image(format="png", width=900, height=500)
        assert png.startswith(b"\x89PNG\r\n\x1a\n")
        assert len(png) > 5_000


def test_excel_export_is_readable_and_keeps_verification_field() -> None:
        workbook = excel_export_bytes({
            "ASSUMPTIONS": pd.DataFrame({
                "assumption_id": ["A1"], "value": [2.5], "source": ["MY ASSUMPTION"],
                "VERIFIED_AGAINST_CASE_TEXT": ["N"],
            }),
            "APPROVALS": pd.DataFrame({
                "risk_id": ["R1"], "approver_role": ["COO"],
                "timestamp": [pd.Timestamp("2026-04-01 10:00")],
                "execution_status": ["NOT_EXECUTED_HUMAN_DECISION_ONLY"],
            }),
        })
        assert len(workbook) > 2_000
        loaded = load_workbook(BytesIO(workbook), read_only=True)
        assert set(loaded.sheetnames) == {"ASSUMPTIONS", "APPROVALS"}
        headers = [cell.value for cell in next(loaded["ASSUMPTIONS"].iter_rows(min_row=1, max_row=1))]
        assert "VERIFIED_AGAINST_CASE_TEXT" in headers
